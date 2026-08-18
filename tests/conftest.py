"""Fixtures partagees.

Tous les chemins du projet derivent de `path_manager.BUREAU`. Rediriger cette
seule variable vers un dossier temporaire suffit donc a isoler completement
les tests du poste de l'utilisateur.
"""

from __future__ import annotations

import os
import zipfile
from xml.sax.saxutils import escape, quoteattr

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from modop import path_manager

LOT = "Lot7"
INSEE = "45001"


@pytest.fixture
def bureau(tmp_path, monkeypatch):
    """Redirige le Bureau vers un dossier temporaire."""
    fake = tmp_path / "Bureau"
    fake.mkdir()
    monkeypatch.setattr(path_manager, "BUREAU", str(fake))
    return fake


@pytest.fixture
def excel_source(bureau):
    """Fichier Excel d'audit realiste, avec ses pieges habituels.

    Contient : types mixtes ("07" et 7), une ligne entierement vide, des
    lignes fantomes porteuses d'une mise en forme sans valeur, et une ligne
    stylee servant a verifier que le style suit la donnee.
    """
    from modop.path_manager import _excel_source_file_path

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Audit"
    worksheet.append(["id", "zone 1", "zone 2", "sens", "adresse"])

    rows = [
        (1, 2, "B", "Nord", "rue A"),
        (2, 1, "A", "Sud", "rue B"),
        (3, 10, "A", "Nord", "rue C"),
        (4, 1, "B", "Nord", "rue D"),
        (5, "07", "A", "Sud", "rue E"),
        (6, 1, "A", "Nord", "rue F"),
        (7, None, "C", "Est", "rue G"),
    ]
    for row in rows:
        worksheet.append(list(row))

    worksheet.append([None] * 5)               # ligne entierement vide
    worksheet.append([8, 3, "C", "Sud", "rue H"])

    # La ligne id=1 est stylee : on verifiera qu'elle garde son style.
    for column in range(1, 6):
        worksheet.cell(row=2, column=column).font = Font(bold=True, color="FFFF0000")

    # Lignes fantomes : mise en forme seule, sans valeur.
    fill = PatternFill("solid", fgColor="FFFFFFCC")
    for row_index in range(30, 80):
        worksheet.cell(row=row_index, column=1).fill = fill

    path = _excel_source_file_path(LOT, INSEE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)
    return path


def _qgs_xml(layers) -> bytes:
    """Genere un .qgs minimal mais conforme a la structure QGIS."""
    tree_entries = []
    layer_entries = []

    for layer in layers:
        # Les URI de source contiennent des & : sans echappement, le XML
        # produit est invalide.
        tree_entries.append(
            f'<layer-tree-layer id={quoteattr(layer["id"])} '
            f'name={quoteattr(layer["name"])} '
            f'source={quoteattr(layer["source"])} '
            f'providerKey={quoteattr(layer["provider"])}/>'
        )

        extent = layer.get("extent")
        extent_xml = ""
        if extent:
            extent_xml = (
                "<extent>"
                f"<xmin>{extent[0]}</xmin><ymin>{extent[1]}</ymin>"
                f"<xmax>{extent[2]}</xmax><ymax>{extent[3]}</ymax>"
                "</extent>"
            )

        layer_entries.append(
            "<maplayer>"
            f"<id>{escape(layer['id'])}</id>"
            f"<datasource>{escape(layer['source'])}</datasource>"
            f"<layername>{escape(layer['name'])}</layername>"
            f"<provider>{escape(layer['provider'])}</provider>"
            f"{extent_xml}"
            "</maplayer>"
        )

    xml = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<qgis projectname="carte_audit_maillage" version="3.34.0">'
        "<title>carte_audit_maillage</title>"
        f'<layer-tree-group>{"".join(tree_entries)}</layer-tree-group>'
        '<mapcanvas name="theMapCanvas"><units>meters</units>'
        "<extent><xmin>0</xmin><ymin>0</ymin><xmax>1</xmax><ymax>1</ymax></extent>"
        "</mapcanvas>"
        f'<projectlayers>{"".join(layer_entries)}</projectlayers>'
        "</qgis>"
    )
    return xml.encode("utf-8")


@pytest.fixture
def default_layers():
    """Deux couches : une commune avec emprise, une couche adresses."""
    return [
        {
            "id": "communes_001",
            "name": "communes",
            "source": "./communes.gpkg",
            "provider": "ogr",
            "extent": (600000, 6700000, 610000, 6710000),
        },
        {
            "id": "adresses_001",
            "name": "adresses_45001",
            "source": "./adresses.csv",
            "provider": "delimitedtext",
            "extent": (601000, 6701000, 609000, 6709000),
        },
    ]


@pytest.fixture
def make_qgz(tmp_path):
    """Fabrique un .qgz a la demande.

    Usage : make_qgz("projet.qgz", layers, extra={"projet.qgd": b"..."})
    """
    def _make(name: str, layers, extra: dict[str, bytes] | None = None,
              directory=None) -> str:
        target_dir = directory or tmp_path
        os.makedirs(target_dir, exist_ok=True)
        path = os.path.join(str(target_dir), name)
        stem = os.path.splitext(name)[0]

        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(f"{stem}.qgs", _qgs_xml(layers))
            for entry_name, content in (extra or {}).items():
                archive.writestr(entry_name, content)

        return path

    return _make


@pytest.fixture
def commune_files(bureau, default_layers, make_qgz):
    """Arborescence complete d'une commune prete a etre traitee.

    Cree LIVRABLE/Lot7/QGIS/45001 avec deux .csv, un .gpkg, et le projet
    modele dans le workspace, au nom que designe path_manager.
    """
    from modop.path_manager import _insee_qgis_path, _qgis_project_path, _workspace_path

    source_dir = _insee_qgis_path(LOT, INSEE)
    os.makedirs(source_dir, exist_ok=True)

    for name, content in (
        ("adresses.csv", "id;zone\n1;A\n"),
        ("maillage.csv", "id;sens\n1;Nord\n"),
        ("communes.gpkg", "donnees binaires factices"),
        ("notes.txt", "fichier non livrable"),
    ):
        with open(os.path.join(source_dir, name), "w", encoding="utf-8") as handle:
            handle.write(content)

    workspace = _workspace_path()
    os.makedirs(workspace, exist_ok=True)

    # Le nom du projet est derive de path_manager : renommer le modele dans
    # le code de production ne casse pas les tests.
    project_path = _qgis_project_path()
    make_qgz(
        os.path.basename(project_path),
        default_layers,
        directory=os.path.dirname(project_path),
    )

    return {
        "source_dir": source_dir,
        "workspace": workspace,
        "project": project_path,
    }


@pytest.fixture
def lot_communes(bureau, default_layers, make_qgz):
    """Lot contenant deux communes aux fichiers differents.

    45001 possede un fichier qui n'existe pas chez 45002 : il sert a detecter
    une contamination du repertoire de travail entre deux communes.
    """
    from modop.path_manager import _insee_qgis_path, _qgis_project_path, _workspace_path

    contenus = {
        "45001": {
            "adresses.csv": "id;zone\n1;A\n",
            "maillage.csv": "id;sens\n1;Nord\n",
            "specifique_45001.csv": "id\n1\n",
            "communes.gpkg": "donnees 45001",
        },
        "45002": {
            "adresses.csv": "id;zone\n2;B\n",
            "maillage.csv": "id;sens\n2;Sud\n",
            "communes.gpkg": "donnees 45002",
        },
    }

    for insee, fichiers in contenus.items():
        source_dir = _insee_qgis_path(LOT, insee)
        os.makedirs(source_dir, exist_ok=True)
        for name, content in fichiers.items():
            with open(os.path.join(source_dir, name), "w", encoding="utf-8") as handle:
                handle.write(content)

    project_path = _qgis_project_path()
    os.makedirs(_workspace_path(), exist_ok=True)
    make_qgz(
        os.path.basename(project_path),
        default_layers,
        directory=os.path.dirname(project_path),
    )

    return {"codes": ["45001", "45002"], "contenus": contenus, "project": project_path}