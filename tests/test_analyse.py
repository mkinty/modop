"""Tests de la génération des PPT vierges (dossier Analyse).

Couvre `modop.services.analyse` : complétion de la colonne "Lien vers le
.ppt" et génération conditionnelle des PPT (case à cocher, template).
"""

from __future__ import annotations

import os
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from modop.constants import PPT_SHAREPOINT_BASE_URL
from modop.path_manager import _insee_analyse_path
from modop.services.analyse import (
    PptTemplateError,
    build_ppt_link,
    build_ppt_sharepoint_link,
    generer_ppts_commune,
    ppt_name_for,
    valider_template_ppt,
)


def _make_pptx_template(path) -> None:
    """Template PPT minimal mais valide : un seul slide avec une variable."""
    slide_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody>"
        "<a:p><a:r><a:t>Adresse : Adresse</a:t></a:r></a:p>"
        "<a:p><a:r><a:t>Commune INSEE</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)


def _make_audit_excel(path, rows) -> None:
    """Excel factice avec les colonnes 'ID erreur' et 'Lien vers le .ppt'."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["ID erreur", "Adresse", "INSEE", "Lien vers le .ppt"])
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


# ---------------------------------------------------------------------------
# Case décochée : rien n'est fait, aucun blocage
# ---------------------------------------------------------------------------

def test_rien_fait_si_case_decochee(tmp_path, bureau):
    """Décochée (défaut), le fichier Excel n'est pas touché du tout."""
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["E1", "1 rue A", "45001", None]])

    resultat = generer_ppts_commune(
        str(excel), "45001", None, generate_ppts=False, verbose=False,
    )

    assert resultat.total == 0
    assert resultat.liens_completes == 0
    assert resultat.ppt_generes == 0
    assert not resultat.avertissements

    worksheet = load_workbook(excel).active
    assert worksheet.cell(2, 4).value is None


def test_case_decochee_ignore_le_template_invalide(tmp_path, bureau):
    """Décochée, un template absent ou invalide ne bloque rien."""
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["E1", "1 rue A", "45001", None]])

    resultat = generer_ppts_commune(
        str(excel), "45001", "/chemin/inexistant.pptx",
        generate_ppts=False, verbose=False,
    )

    assert resultat.ppt_generes == 0
    assert not resultat.avertissements


# ---------------------------------------------------------------------------
# Case cochée sans template valide : blocage
# ---------------------------------------------------------------------------

def test_template_absent_bloque_lexecution(tmp_path, bureau):
    """Case cochée, aucun template configuré : PptTemplateError levée."""
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["E1", "1 rue A", "45001", None]])

    with pytest.raises(PptTemplateError):
        generer_ppts_commune(
            str(excel), "45001", None, generate_ppts=True, verbose=False,
        )

    # Rien n'a ete ecrit : le blocage intervient avant l'ouverture du fichier.
    worksheet = load_workbook(excel).active
    assert worksheet.cell(2, 4).value is None


def test_template_introuvable_bloque_lexecution(tmp_path, bureau):
    """Case cochée, chemin de template configuré mais introuvable : bloque."""
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["E1", "1 rue A", "45001", None]])

    with pytest.raises(PptTemplateError):
        generer_ppts_commune(
            str(excel), "45001", "/chemin/inexistant.pptx",
            generate_ppts=True, verbose=False,
        )


def test_blocage_independant_de_lexcel(tmp_path, bureau):
    """Le blocage sur template invalide intervient meme sans audit a traiter."""
    with pytest.raises(PptTemplateError):
        generer_ppts_commune(
            str(tmp_path / "absent.xlsx"), "45001", None,
            generate_ppts=True, verbose=False,
        )


# ---------------------------------------------------------------------------
# Génération effective des PPT, avec un template valide
# ---------------------------------------------------------------------------

def test_lignes_sans_id_ignorees(tmp_path, bureau):
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["", "1 rue A", "45001", None], ["E2", "2 rue B", "45001", None]])
    template = tmp_path / "template.pptx"
    _make_pptx_template(template)

    resultat = generer_ppts_commune(
        str(excel), "45001", str(template), generate_ppts=True, verbose=False,
    )

    assert resultat.total == 1
    assert resultat.liens_completes == 1


def test_ppt_genere_si_case_cochee(tmp_path, bureau):
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [["E1", "1 rue A", "45001", None]])
    template = tmp_path / "template.pptx"
    _make_pptx_template(template)

    resultat = generer_ppts_commune(
        str(excel), "45001", str(template), generate_ppts=True, verbose=False,
    )

    assert resultat.ppt_generes == 1
    assert resultat.erreurs == 0

    worksheet = load_workbook(excel).active
    lien = worksheet.cell(2, 4).value
    # La cellule porte le lien SharePoint (ouverture navigateur), pas le
    # chemin local ; le fichier réellement généré se vérifie via
    # build_ppt_link.
    assert lien == build_ppt_sharepoint_link("45001", "E1")
    assert lien.startswith(PPT_SHAREPOINT_BASE_URL + "/")
    assert lien.endswith("/" + ppt_name_for("E1") + "?web=1")
    assert "\\" not in lien
    assert os.path.isfile(build_ppt_link("45001", "E1"))


def test_plusieurs_lignes_generation_parallele(tmp_path, bureau):
    excel = tmp_path / "audit_45001.xlsx"
    _make_audit_excel(excel, [
        ["E1", "1 rue A", "45001", None],
        ["E2", "2 rue B", "45001", None],
        ["E3", "3 rue C", "45001", None],
    ])
    template = tmp_path / "template.pptx"
    _make_pptx_template(template)

    resultat = generer_ppts_commune(
        str(excel), "45001", str(template), generate_ppts=True, verbose=False,
    )

    assert resultat.total == 3
    assert resultat.liens_completes == 3
    assert resultat.ppt_generes == 3
    for id_err in ("E1", "E2", "E3"):
        assert os.path.isfile(build_ppt_link("45001", id_err))


# ---------------------------------------------------------------------------
# Cas limites : template valide mais donnees absentes/incompletes
# ---------------------------------------------------------------------------

def test_colonnes_absentes_est_un_simple_avertissement(tmp_path, bureau):
    """Un audit sans les colonnes attendues n'interrompt pas le traitement,
    des lors qu'un template valide est fourni (donc pas de blocage)."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["zone 1", "zone 2", "sens"])
    worksheet.append([1, "A", "Nord"])
    excel = tmp_path / "audit_45002.xlsx"
    workbook.save(excel)

    template = tmp_path / "template.pptx"
    _make_pptx_template(template)

    resultat = generer_ppts_commune(
        str(excel), "45002", str(template), generate_ppts=True, verbose=False,
    )

    assert resultat.total == 0
    assert resultat.liens_completes == 0
    assert resultat.avertissements


def test_fichier_excel_absent_avec_template_valide(tmp_path, bureau):
    """Le fichier Excel absent reste une simple anomalie de donnees, non
    bloquante, des lors que le template est valide."""
    template = tmp_path / "template.pptx"
    _make_pptx_template(template)

    resultat = generer_ppts_commune(
        str(tmp_path / "absent.xlsx"), "45001", str(template),
        generate_ppts=True, verbose=False,
    )
    assert resultat.avertissements
    assert resultat.total == 0


# ---------------------------------------------------------------------------
# valider_template_ppt : le contrôle isolé, réutilisable en amont d'un lot
# ---------------------------------------------------------------------------

def test_valider_template_ppt_chemin_vide():
    with pytest.raises(PptTemplateError):
        valider_template_ppt(None)
    with pytest.raises(PptTemplateError):
        valider_template_ppt("")


def test_valider_template_ppt_fichier_introuvable(tmp_path):
    with pytest.raises(PptTemplateError):
        valider_template_ppt(str(tmp_path / "absent.pptx"))


def test_valider_template_ppt_valide(tmp_path):
    template = tmp_path / "template.pptx"
    _make_pptx_template(template)
    # Ne doit lever aucune exception.
    valider_template_ppt(str(template))


# ---------------------------------------------------------------------------
# Nommage et chemins
# ---------------------------------------------------------------------------

def test_ppt_name_for():
    assert ppt_name_for("E1") == "cas_analyse_E1.pptx"


def test_build_ppt_link_dans_le_dossier_analyse(bureau):
    lien = build_ppt_link("45001", "E1")
    assert lien == os.path.join(_insee_analyse_path("45001"), "cas_analyse_E1.pptx")
