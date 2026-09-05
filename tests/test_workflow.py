"""Tests d'integration du module modop.services.workflow.

Ces tests parcourent les cinq etapes du mode operatoire de bout en bout.
"""

from __future__ import annotations

import os
import zipfile

import pytest

from modop.path_manager import (
    _deliverable_archive_path,
    _qgis_project_path,
    _qgis_saved_project_path,
    _workspace_path,
)
from modop.services.qgis_project import QgisProject
from modop.services.workflow import prepare_commune_deliverable
from tests.conftest import INSEE, LOT


def test_le_modele_suit_path_manager(commune_files):
    """Garde-fou : le fixture doit fabriquer le projet la ou path_manager
    l'attend. Renommer le modele dans path_manager ne doit pas faire tomber
    toute la suite avec un FileNotFoundError peu parlant."""
    assert commune_files["project"] == _qgis_project_path()
    assert os.path.isfile(_qgis_project_path())


def test_livrable_complet(commune_files):
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert resultat.ok
    assert resultat.insee == INSEE
    assert len(resultat.copied_files) == 4
    assert resultat.archive_file == _deliverable_archive_path(INSEE)


def test_projet_enregistre_selon_la_convention(commune_files):
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)
    nom = os.path.basename(resultat.project_file)

    assert resultat.project_file == _qgis_saved_project_path(INSEE)
    assert os.path.isfile(resultat.project_file)

    # Le separateur est fixe par path_manager ; ce qui est verifie ici, c'est
    # le contrat : le nom porte le code INSEE et l'extension du projet.
    assert nom.startswith("carte_audit")
    assert INSEE in nom
    assert nom.endswith(".qgz")


def test_carte_centree_sur_la_commune(commune_files):
    """L'emprise du canevas doit correspondre a celle de la couche communes."""
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)
    project = QgisProject.open(resultat.project_file)

    emprise = project.canvas_extent
    assert emprise is not None
    assert emprise.xmin < 600000 < emprise.xmax        # marge appliquee
    assert emprise.ymin < 6700000 < emprise.ymax


def test_couche_de_zoom_detectee_automatiquement(commune_files):
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)
    assert resultat.zoom_layer == "communes"


def test_couche_de_zoom_imposee(commune_files):
    resultat = prepare_commune_deliverable(
        LOT, INSEE, zoom_layer="adresses_45001", verbose=False
    )
    assert resultat.zoom_layer == "adresses_45001"


def test_archive_contient_csv_et_projet(commune_files):
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    with zipfile.ZipFile(resultat.archive_file) as archive:
        noms = sorted(archive.namelist())

    projet = os.path.basename(_qgis_saved_project_path(INSEE))
    assert noms == sorted(["adresses.csv", "maillage.csv", projet])


def test_fichiers_copies_dans_le_workspace(commune_files):
    prepare_commune_deliverable(LOT, INSEE, verbose=False)
    workspace = _workspace_path()

    assert os.path.isfile(os.path.join(workspace, "adresses.csv"))
    assert os.path.isfile(os.path.join(workspace, "communes.gpkg"))


def test_projet_source_non_modifie(commune_files):
    """carte_audit_maillage.qgz reste le modele reutilisable."""
    source = commune_files["project"]
    prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert QgisProject.open(source).title == "carte_audit_maillage"


def test_aucune_anomalie_quand_les_donnees_sont_copiees(commune_files):
    """La copie prealable rend les sources relatives du projet valides."""
    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)
    assert resultat.warnings == []


def test_source_manquante_remontee_en_avertissement(commune_files):
    """Une couche dont la donnee n'a pas ete livree est signalee."""
    os.remove(os.path.join(commune_files["source_dir"], "communes.gpkg"))
    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)

    assert any("Source introuvable" in message for message in resultat.warnings)
    assert resultat.ok      # sans strict, le traitement va jusqu'au bout


def test_mode_strict_interrompt_le_traitement(commune_files):
    os.remove(os.path.join(commune_files["source_dir"], "communes.gpkg"))

    with pytest.raises(ValueError, match="Source introuvable"):
        prepare_commune_deliverable(
            LOT, INSEE, sort_excel=False, strict=True, verbose=False
        )


def test_commune_absente(bureau):
    with pytest.raises(FileNotFoundError):
        prepare_commune_deliverable(LOT, "99999", verbose=False)


def test_projet_qgis_absent(commune_files):
    os.remove(commune_files["project"])

    with pytest.raises(FileNotFoundError, match="Projet QGIS"):
        prepare_commune_deliverable(LOT, INSEE, verbose=False)


def test_couche_de_zoom_inexistante(commune_files):
    with pytest.raises(LookupError, match="parcelles"):
        prepare_commune_deliverable(LOT, INSEE, zoom_layer="parcelles", verbose=False)


def test_projet_sans_couche_exploitable(commune_files, make_qgz):
    """Aucune couche n'a d'emprise : le centrage ne peut pas etre determine."""
    projet = make_qgz("modele.qgz", [
        {"id": "c1", "name": "vide", "provider": "ogr", "source": "./x.gpkg", "extent": None}
    ])

    with pytest.raises(LookupError, match="emprise"):
        prepare_commune_deliverable(LOT, INSEE, project_path=projet, verbose=False)


def test_traitement_reproductible(commune_files):
    """Relancer la commune produit le meme livrable."""
    premier = prepare_commune_deliverable(LOT, INSEE, verbose=False)
    with zipfile.ZipFile(premier.archive_file) as archive:
        avant = sorted(archive.namelist())

    second = prepare_commune_deliverable(LOT, INSEE, verbose=False)
    with zipfile.ZipFile(second.archive_file) as archive:
        apres = sorted(archive.namelist())

    assert avant == apres
    assert second.ok


def test_journalisation(commune_files, capsys):
    prepare_commune_deliverable(LOT, INSEE, verbose=True)
    sortie = capsys.readouterr().out

    assert "[copie]" in sortie
    assert "[qgis]" in sortie
    assert "[zip]" in sortie
    assert "[ok]" in sortie


# ---------------------------------------------------------------------------
# Integration du tri Excel
# ---------------------------------------------------------------------------

def test_fichier_audit_trie_et_livre(commune_files, excel_source):
    from openpyxl import load_workbook
    from modop.path_manager import _excel_destination_file_path

    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert resultat.complete
    assert resultat.excel_file == _excel_destination_file_path(INSEE)

    worksheet = load_workbook(resultat.excel_file).active
    zones = [row[1] for row in worksheet.iter_rows(min_row=2, values_only=True)
             if any(v is not None for v in row)]
    assert zones == [1, 1, 1, 2, 3, "07", 10, None]


def test_audit_livre_hors_archive(commune_files, excel_source):
    """Le fichier Excel est livre dans le dossier commune, pas dans le ZIP."""
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    with zipfile.ZipFile(resultat.archive_file) as archive:
        assert not any(nom.endswith(".xlsx") for nom in archive.namelist())


def test_audit_absent_nempeche_pas_le_livrable(commune_files):
    """Les deux chaines sont independantes."""
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert resultat.ok
    assert not resultat.complete
    assert resultat.excel_file == ""
    assert any("audit absent" in message for message in resultat.warnings)


def test_audit_absent_bloque_en_mode_strict(commune_files):
    with pytest.raises(ValueError, match="audit absent"):
        prepare_commune_deliverable(LOT, INSEE, strict=True, verbose=False)


def test_echec_du_tri_signale(commune_files, bureau):
    """Un fichier d'audit sans les colonnes attendues est signale."""
    from openpyxl import Workbook
    from modop.path_manager import _excel_source_file_path

    workbook = Workbook()
    workbook.active.append(["id", "zone 1", "sens"])   # zone 2 manquante
    path = _excel_source_file_path(LOT, INSEE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)

    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert resultat.ok
    assert any("echoue" in message for message in resultat.warnings)


def test_sort_excel_desactive(commune_files, excel_source):
    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)

    assert resultat.ok
    assert resultat.excel_file == ""
    assert resultat.warnings == []


def test_source_audit_intacte(commune_files, excel_source):
    from openpyxl import load_workbook

    avant = [row[0] for row in load_workbook(excel_source).active.iter_rows(
        min_row=2, values_only=True)]

    prepare_commune_deliverable(LOT, INSEE, verbose=False)

    apres = [row[0] for row in load_workbook(excel_source).active.iter_rows(
        min_row=2, values_only=True)]
    assert avant == apres


def test_journalisation_complete(commune_files, excel_source, capsys):
    prepare_commune_deliverable(LOT, INSEE, verbose=True)
    sortie = capsys.readouterr().out

    for etape in ("[excel]", "[copie]", "[qgis]", "[zip]", "[ok]"):
        assert etape in sortie

    # L'ordre du mode operatoire est respecte.
    assert sortie.index("[excel]") < sortie.index("[copie]") < sortie.index("[ok]")


# ---------------------------------------------------------------------------
# Arborescence de la commune
# ---------------------------------------------------------------------------

def test_dossiers_carte_et_analyse_crees(commune_files):
    from modop.path_manager import _insee_analyse_path, _insee_carte_path

    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)

    assert resultat.carte_dir == _insee_carte_path(INSEE)
    assert resultat.analysis_dir == _insee_analyse_path(INSEE)
    assert os.path.isdir(resultat.carte_dir)
    assert os.path.isdir(resultat.analysis_dir)
    assert os.path.basename(resultat.analysis_dir) == "Analyse"


def test_analyse_est_un_sous_dossier_de_la_commune(commune_files):
    from modop.path_manager import _insee_dep_path

    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)
    assert os.path.dirname(resultat.analysis_dir) == _insee_dep_path(INSEE)


def test_analyse_reste_vide(commune_files):
    """Le dossier est préparé, pas rempli : rien du livrable n'y va."""
    resultat = prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)
    assert os.listdir(resultat.analysis_dir) == []


def test_dossiers_crees_meme_sans_audit(commune_files):
    """L'arborescence précède le tri : elle existe même si l'audit manque."""
    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert os.path.isdir(resultat.analysis_dir)
    assert any("audit absent" in message for message in resultat.warnings)


def test_dossier_existant_non_ecrase(commune_files):
    """Un dossier Analyse déjà rempli doit être conservé."""
    from modop.path_manager import _insee_analyse_path

    depart = _insee_analyse_path(INSEE)
    with open(os.path.join(depart, "note.txt"), "w", encoding="utf-8") as handle:
        handle.write("travail en cours")

    prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=False)
    assert os.listdir(depart) == ["note.txt"]


def test_journal_signale_les_dossiers(commune_files, capsys):
    prepare_commune_deliverable(LOT, INSEE, sort_excel=False, verbose=True)
    assert "[dossier]" in capsys.readouterr().out


# ---------------------------------------------------------------------------
# PPT vierges du dossier Analyse
# ---------------------------------------------------------------------------

@pytest.fixture
def excel_source_ppt(bureau):
    """Audit realiste, complet cette fois des colonnes utilisees pour les PPT."""
    import zipfile as _zipfile
    from openpyxl import Workbook
    from modop.path_manager import _excel_source_file_path

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["zone 1", "zone 2", "sens", "ID erreur", "Adresse",
                      "INSEE", "Lien vers le .ppt"])
    worksheet.append([2, "B", "Nord", "E1", "1 rue A", INSEE, None])
    worksheet.append([1, "A", "Sud", "E2", "2 rue B", INSEE, None])

    path = _excel_source_file_path(LOT, INSEE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)
    return path


@pytest.fixture
def pptx_template(tmp_path):
    """Template PPT minimal mais valide, pour la generation."""
    import zipfile as _zipfile

    slide_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody>"
        "<a:p><a:r><a:t>Adresse : Adresse</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    path = tmp_path / "template.pptx"
    with _zipfile.ZipFile(path, "w", _zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
    return str(path)


def test_rien_ne_change_par_defaut(commune_files, excel_source_ppt):
    """La case n'etant pas cochee par defaut, l'audit n'est pas touche."""
    from openpyxl import load_workbook

    resultat = prepare_commune_deliverable(LOT, INSEE, verbose=False)

    assert resultat.ppt_liens == 0
    assert resultat.ppt_generes == 0
    assert os.listdir(resultat.analysis_dir) == []

    worksheet = load_workbook(resultat.excel_file).active
    entetes = {worksheet.cell(1, c).value: c for c in range(1, worksheet.max_column + 1)}
    col_lien = entetes["Lien vers le .ppt"]
    liens = [worksheet.cell(r, col_lien).value for r in (2, 3)]
    assert liens == [None, None]


def test_case_cochee_sans_template_bloque_le_traitement(commune_files, excel_source_ppt):
    """Case cochée sans template valide : le traitement de la commune
    s'interrompt (PptTemplateError), il ne se contente pas d'un avertissement."""
    from modop.services.analyse import PptTemplateError

    with pytest.raises(PptTemplateError):
        prepare_commune_deliverable(LOT, INSEE, generate_ppts=True, verbose=False)

    # Le blocage intervient avant l'ecriture des liens : l'audit n'est pas modifie.
    from openpyxl import load_workbook

    worksheet = load_workbook(excel_source_ppt).active
    entetes = {worksheet.cell(1, c).value: c for c in range(1, worksheet.max_column + 1)}
    col_lien = entetes["Lien vers le .ppt"]
    assert worksheet.cell(2, col_lien).value is None


def test_ppt_generes_si_case_cochee(commune_files, excel_source_ppt, pptx_template):
    resultat = prepare_commune_deliverable(
        LOT, INSEE, pptx_template_path=pptx_template, generate_ppts=True,
        verbose=False,
    )

    assert resultat.ppt_liens == 2
    assert resultat.ppt_generes == 2
    fichiers = os.listdir(resultat.analysis_dir)
    assert len(fichiers) == 2
    assert all(nom.startswith("cas_analyse_") and nom.endswith(".pptx")
               for nom in fichiers)


def test_journal_signale_le_blocage_sur_template(commune_files, excel_source_ppt, capsys):
    from modop.services.analyse import PptTemplateError

    with pytest.raises(PptTemplateError):
        prepare_commune_deliverable(LOT, INSEE, generate_ppts=True, verbose=True)

    assert "[ppt] ❌" in capsys.readouterr().out


def test_journal_signale_l_etape_ppt(commune_files, excel_source_ppt, pptx_template, capsys):
    prepare_commune_deliverable(
        LOT, INSEE, pptx_template_path=pptx_template, generate_ppts=True, verbose=True,
    )
    assert "[ppt]" in capsys.readouterr().out


def test_journal_muet_sur_ppt_si_case_decochee(commune_files, excel_source_ppt, capsys):
    prepare_commune_deliverable(LOT, INSEE, verbose=True)
    assert "[ppt]" not in capsys.readouterr().out