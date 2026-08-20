"""Tests de la preservation des extensions xlsx.

openpyxl supprime les blocs <extLst> des feuilles : listes deroulantes dont la
source est sur une autre feuille, mises en forme conditionnelles avancees,
plages protegees. Ces tests verifient qu'elles survivent au tri.
"""

from __future__ import annotations

import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from modop.path_manager import _excel_source_file_path
from modop.services.excel import format_excel_file
from modop.services.xlsx_extensions import (
    describe_extensions,
    find_extensions,
    restore_extensions,
)
from tests.conftest import INSEE, LOT

URI_VALIDATION = "{CCE6A557-97BC-4B89-ADB6-D9C93CAAB3DF}"
URI_FORMAT = "{78C0D931-6437-407D-A8EE-F0AAD7539E65}"

#: Validation x14 : liste deroulante dont la source est sur une autre feuille.
#: C'est le cas courant qu'openpyxl ne sait pas conserver.
EXT_VALIDATION = (
    f'<extLst><ext uri="{URI_VALIDATION}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" '
    'count="1"><x14:dataValidation type="list" allowBlank="1">'
    "<x14:formula1><xm:f>Listes!$A$1:$A$3</xm:f></x14:formula1>"
    "<xm:sqref>E2:E100</xm:sqref>"
    "</x14:dataValidation></x14:dataValidations></ext></extLst>"
)


@pytest.fixture
def classeur_avec_extension(tmp_path):
    """Fabrique un .xlsx portant une validation x14."""
    def _make(extension: str = EXT_VALIDATION, nom: str = "avec_ext.xlsx",
              racine_ns: bool = False) -> str:
        base = tmp_path / "base.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["id", "zone 1", "zone 2", "sens", "statut"])
        worksheet.append([1, 2, "B", "Nord", "OK"])
        worksheet.append([2, 1, "A", "Sud", "NOK"])
        workbook.save(base)

        cible = tmp_path / nom
        with zipfile.ZipFile(base) as source, \
             zipfile.ZipFile(cible, "w", zipfile.ZIP_DEFLATED) as sortie:
            for item in source.infolist():
                contenu = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    xml = contenu.decode("utf-8")
                    if racine_ns:
                        # Cas reel : le prefixe est declare sur <worksheet>,
                        # pas dans le fragment lui-meme.
                        xml = xml.replace(
                            "<worksheet ",
                            '<worksheet xmlns:x14="http://schemas.microsoft.com/'
                            'office/spreadsheetml/2009/9/main" ', 1)
                        extension_locale = extension.replace(
                            ' xmlns:x14="http://schemas.microsoft.com/office/'
                            'spreadsheetml/2009/9/main"', "")
                    else:
                        extension_locale = extension
                    xml = xml.replace("</worksheet>", extension_locale + "</worksheet>")
                    contenu = xml.encode("utf-8")
                sortie.writestr(item, contenu)

        return str(cible)

    return _make


# ---------------------------------------------------------------------------
# Relevé
# ---------------------------------------------------------------------------

def test_find_extensions(classeur_avec_extension):
    trouves = find_extensions(classeur_avec_extension())

    assert list(trouves) == ["xl/worksheets/sheet1.xml"]
    assert URI_VALIDATION in trouves["xl/worksheets/sheet1.xml"]


def test_find_extensions_classeur_ordinaire(tmp_path):
    chemin = tmp_path / "simple.xlsx"
    Workbook().save(chemin)

    assert find_extensions(str(chemin)) == {}


def test_find_extensions_fichier_illisible(tmp_path):
    chemin = tmp_path / "casse.xlsx"
    chemin.write_text("ceci n'est pas un zip", encoding="utf-8")

    assert find_extensions(str(chemin)) == {}


def test_find_extensions_fichier_absent(tmp_path):
    assert find_extensions(str(tmp_path / "absent.xlsx")) == {}


def test_espace_de_noms_rappele(classeur_avec_extension):
    """Un prefixe declare sur <worksheet> doit suivre le fragment deplace."""
    trouves = find_extensions(classeur_avec_extension(racine_ns=True))
    fragment = trouves["xl/worksheets/sheet1.xml"]

    assert 'xmlns:x14="' in fragment


def test_describe_extensions(classeur_avec_extension):
    trouves = find_extensions(classeur_avec_extension())
    assert describe_extensions(trouves) == ["listes déroulantes"]


def test_describe_extension_inconnue():
    fragment = {"s": '<extLst><ext uri="{00000000-0000-0000-0000-000000000000}"/></extLst>'}
    assert describe_extensions(fragment) == ["extension non identifiée"]


def test_describe_sans_doublon():
    fragment = {
        "a": f'<extLst><ext uri="{URI_VALIDATION}"/></extLst>',
        "b": f'<extLst><ext uri="{URI_VALIDATION}"/></extLst>',
    }
    assert describe_extensions(fragment) == ["listes déroulantes"]


# ---------------------------------------------------------------------------
# Réinjection
# ---------------------------------------------------------------------------

def test_restore_extensions(classeur_avec_extension, tmp_path):
    source = classeur_avec_extension()
    trouves = find_extensions(source)

    # Passage par openpyxl : l'extension est perdue.
    depouille = str(tmp_path / "depouille.xlsx")
    load_workbook(source).save(depouille)
    assert find_extensions(depouille) == {}

    assert restore_extensions(depouille, trouves)
    assert URI_VALIDATION in find_extensions(depouille)["xl/worksheets/sheet1.xml"]


def test_restore_conserve_un_fichier_lisible(classeur_avec_extension, tmp_path):
    source = classeur_avec_extension()
    trouves = find_extensions(source)
    cible = str(tmp_path / "cible.xlsx")
    load_workbook(source).save(cible)
    restore_extensions(cible, trouves)

    worksheet = load_workbook(cible).active
    assert [c.value for c in worksheet[1]] == ["id", "zone 1", "zone 2", "sens", "statut"]


def test_restore_sans_extension(tmp_path):
    chemin = str(tmp_path / "simple.xlsx")
    Workbook().save(chemin)

    assert restore_extensions(chemin, {}) is False


def test_restore_ne_double_pas_un_extlst_existant(classeur_avec_extension):
    """Un fichier qui possede deja le bloc ne doit pas en recevoir un second."""
    source = classeur_avec_extension()
    trouves = find_extensions(source)

    assert restore_extensions(source, trouves) is False

    xml = zipfile.ZipFile(source).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert xml.count("<extLst>") == 1


def test_restore_laisse_le_fichier_intact_en_cas_d_echec(tmp_path):
    chemin = tmp_path / "casse.xlsx"
    chemin.write_text("pas un zip", encoding="utf-8")

    assert restore_extensions(str(chemin), {"x": "<extLst/>"}) is False
    assert chemin.read_text(encoding="utf-8") == "pas un zip"


def test_restore_ne_laisse_pas_de_fichier_temporaire(classeur_avec_extension, tmp_path):
    source = classeur_avec_extension()
    cible = str(tmp_path / "cible.xlsx")
    load_workbook(source).save(cible)
    restore_extensions(cible, find_extensions(source))

    assert not (tmp_path / "cible.xlsx.ext.tmp").exists()


# ---------------------------------------------------------------------------
# Intégration au tri
# ---------------------------------------------------------------------------

def _installer_source(bureau, classeur: str) -> str:
    import os
    import shutil

    chemin = _excel_source_file_path(LOT, INSEE)
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    shutil.copy(classeur, chemin)
    return chemin


def test_le_tri_preserve_les_listes_deroulantes(bureau, classeur_avec_extension):
    _installer_source(bureau, classeur_avec_extension())
    resultat = format_excel_file(LOT, INSEE, verbose=False)

    xml = zipfile.ZipFile(resultat).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "x14:dataValidation" in xml
    assert "E2:E100" in xml       # la plage visée est conservée


def test_le_tri_reste_correct_avec_extension(bureau, classeur_avec_extension):
    _installer_source(bureau, classeur_avec_extension())
    resultat = format_excel_file(LOT, INSEE, verbose=False)

    worksheet = load_workbook(resultat).active
    assert [r[1] for r in worksheet.iter_rows(min_row=2, values_only=True)] == [1, 2]


def test_keep_extensions_desactive(bureau, classeur_avec_extension):
    _installer_source(bureau, classeur_avec_extension())
    resultat = format_excel_file(LOT, INSEE, keep_extensions=False, verbose=False)

    assert find_extensions(resultat) == {}


def test_avertissement_openpyxl_reformule(bureau, classeur_avec_extension, capsys):
    """L'avertissement brut est capté puis réémis dans le journal."""
    _installer_source(bureau, classeur_avec_extension())
    format_excel_file(LOT, INSEE, verbose=True)

    sortie = capsys.readouterr().out
    assert "openpyxl : Data Validation extension" in sortie
    assert "Extensions restaurées : listes déroulantes" in sortie