"""Tests du module modop.services.excel."""

from __future__ import annotations

import os
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from modop.path_manager import _excel_destination_file_path, _excel_source_file_path
from modop.services import excel
from modop.services.excel import (
    _last_data_row,
    _natural_chunks,
    _normalize_header,
    _sort_key,
    _to_number,
    format_excel_file,
)
from tests.conftest import INSEE, LOT


# ---------------------------------------------------------------------------
# Normalisation des en-tetes
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "brut, attendu",
    [
        ("zone 1", "zone 1"),
        ("  ZONE 1  ", "zone 1"),
        ("Zone\u00a01", "zone 1"),      # espace insecable
        ("zone   2", "zone 2"),         # espaces multiples
        ("Séns", "sens"),               # accent
        (None, ""),
        (42, "42"),                     # en-tete numerique
    ],
)
def test_normalize_header(brut, attendu):
    assert _normalize_header(brut) == attendu


# ---------------------------------------------------------------------------
# Cle de tri
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "texte, attendu",
    [
        ("07", 7.0),
        ("1,5", 1.5),
        ("1 200", 1200.0),
        ("1\u00a0200", 1200.0),
        ("Z2", None),
        ("1.2.3", None),
        ("", None),
        ("nan", None),                  # incomparable : rejete
        ("inf", None),
    ],
)
def test_to_number(texte, attendu):
    assert _to_number(texte) == attendu


def test_natural_chunks_trie_z2_avant_z10():
    assert _natural_chunks("z2") < _natural_chunks("z10")


def test_sort_key_ordre_des_types():
    """Nombres < dates < texte < cellules vides."""
    assert _sort_key(1) < _sort_key(date(2024, 1, 1))
    assert _sort_key(date(2024, 1, 1)) < _sort_key("abc")
    assert _sort_key("abc") < _sort_key(None)


def test_sort_key_texte_numerique_avec_les_nombres():
    """'07' se trie avec 7, donc entre 3 et 10."""
    assert _sort_key(3) < _sort_key("07") < _sort_key(10)


def test_sort_key_bool_traite_comme_nombre():
    assert _sort_key(False) < _sort_key(True) < _sort_key(2)


def test_sort_key_datetime_conserve_l_heure():
    assert _sort_key(datetime(2024, 1, 1, 8)) < _sort_key(datetime(2024, 1, 1, 9))


def test_sort_key_chaine_vide_vaut_cellule_vide():
    assert _sort_key("   ") == _sort_key(None)


def test_sort_key_aucun_typeerror_sur_types_melanges():
    """Le rang de type garantit qu'aucune comparaison ne peut echouer."""
    valeurs = [3, "Z2", None, date(2024, 1, 1), True, "07", "", 2.5]
    sorted(valeurs, key=_sort_key)      # ne doit pas lever


# ---------------------------------------------------------------------------
# Lignes fantomes
# ---------------------------------------------------------------------------

def test_last_data_row_ignore_les_lignes_fantomes(tmp_path):
    from openpyxl.styles import PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["h"])
    worksheet.append(["a"])
    for row_index in range(20, 60):
        worksheet.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor="FFFFFF00")

    assert worksheet.max_row >= 59          # openpyxl compte les fantomes
    assert _last_data_row(worksheet, 2) == 2


def test_last_data_row_sans_donnees():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["h"])
    assert _last_data_row(worksheet, 2) == 1


# ---------------------------------------------------------------------------
# Traitement complet
# ---------------------------------------------------------------------------

def _lire(path):
    """Retourne les lignes de donnees du fichier trie."""
    worksheet = load_workbook(path).active
    return [row for row in worksheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)]


def test_format_excel_file_trie_correctement(excel_source):
    resultat = format_excel_file(LOT, INSEE, verbose=False)

    assert resultat == _excel_destination_file_path(INSEE)
    zones = [row[1] for row in _lire(resultat)]
    assert zones == [1, 1, 1, 2, 3, "07", 10, None]


def test_format_excel_file_tri_sur_les_trois_colonnes(excel_source):
    """A zone 1 egale, zone 2 departage ; puis sens."""
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    trois_premiers = [row[1:4] for row in _lire(resultat)[:3]]

    assert trois_premiers == [
        (1, "A", "Nord"),
        (1, "A", "Sud"),
        (1, "B", "Nord"),
    ]


def test_format_excel_file_ne_vide_pas_le_fichier(excel_source):
    """Regression : l'ecriture sur cellules vivantes vidait le fichier."""
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    lignes = _lire(resultat)

    assert len(lignes) == 8
    # Les identifiants d'origine sont tous presents, aucun doublon.
    assert sorted(row[0] for row in lignes) == [1, 2, 3, 4, 5, 6, 7, 8]


def test_format_excel_file_supprime_les_lignes_vides(excel_source):
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    worksheet = load_workbook(resultat).active

    valeurs = list(worksheet.iter_rows(min_row=2, max_row=9, values_only=True))
    assert all(any(v is not None for v in row) for row in valeurs)


def test_format_excel_file_conserve_le_source(excel_source):
    avant = load_workbook(excel_source).active
    ordre_avant = [row[0] for row in avant.iter_rows(min_row=2, values_only=True)]

    format_excel_file(LOT, INSEE, verbose=False)

    apres = load_workbook(excel_source).active
    ordre_apres = [row[0] for row in apres.iter_rows(min_row=2, values_only=True)]
    assert ordre_avant == ordre_apres


def test_format_excel_file_le_style_suit_la_donnee(excel_source):
    """La ligne id=1 etait en rouge gras : elle doit l'etre encore."""
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    worksheet = load_workbook(resultat).active

    cellule = next(c for c in worksheet["A"] if c.value == 1)
    assert cellule.font.bold is True
    assert cellule.font.color.rgb == "FFFF0000"


def test_format_excel_file_est_idempotent(excel_source):
    """Deux passages consecutifs donnent le meme resultat."""
    premier = _lire(format_excel_file(LOT, INSEE, verbose=False))
    second = _lire(format_excel_file(LOT, INSEE, verbose=False))
    assert premier == second


# ---------------------------------------------------------------------------
# Structures de fichier particulieres
# ---------------------------------------------------------------------------

def _ecrire_source(bureau, lignes, entete_ligne=1, titre=None, feuille_vide=False):
    """Cree un fichier source sur mesure."""
    workbook = Workbook()
    worksheet = workbook.active

    if feuille_vide:
        worksheet.title = "Vide"
        worksheet = workbook.create_sheet("Donnees")

    if titre:
        worksheet["A1"] = titre

    for offset, valeurs in enumerate(lignes):
        for column, valeur in enumerate(valeurs, start=1):
            worksheet.cell(row=entete_ligne + offset, column=column, value=valeur)

    path = _excel_source_file_path(LOT, INSEE)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)
    return path


def test_entete_hors_ligne_1(bureau):
    _ecrire_source(
        bureau,
        [["id", "zone 1", "zone 2", "sens"], [1, 2, "B", "Sud"], [2, 1, "A", "Nord"]],
        entete_ligne=3,
        titre="EXPORT AUDIT",
    )
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    worksheet = load_workbook(resultat).active

    assert worksheet["A4"].value == 2       # la zone 1 la plus petite remonte
    assert worksheet["A1"].value == "EXPORT AUDIT"   # le titre n'a pas bouge


def test_entete_sur_une_autre_feuille(bureau):
    _ecrire_source(
        bureau,
        [["id", "zone 1", "zone 2", "sens"], [1, 2, "B", "Sud"], [2, 1, "A", "Nord"]],
        feuille_vide=True,
    )
    resultat = format_excel_file(LOT, INSEE, verbose=False)

    worksheet = load_workbook(resultat)["Donnees"]
    assert worksheet["A2"].value == 2


def test_colonne_manquante_retourne_none(bureau, capsys):
    _ecrire_source(bureau, [["id", "zone 1", "sens"], [1, 2, "Sud"]])

    assert format_excel_file(LOT, INSEE, verbose=False) is None
    assert "introuvable" in capsys.readouterr().out.lower()


def test_fusion_dans_les_donnees_refusee(bureau, capsys):
    _ecrire_source(
        bureau,
        [["id", "zone 1", "zone 2", "sens"], [1, 2, "B", "Sud"], [2, 1, "A", "Nord"]],
    )
    path = _excel_source_file_path(LOT, INSEE)
    workbook = load_workbook(path)
    workbook.active.merge_cells("A2:B2")
    workbook.save(path)

    assert format_excel_file(LOT, INSEE, verbose=False) is None
    assert "fusionn" in capsys.readouterr().out.lower()


def test_source_absent_retourne_none(bureau, capsys):
    assert format_excel_file(LOT, "99999", verbose=False) is None
    assert "introuvable" in capsys.readouterr().out.lower()


def test_aucune_donnee_sous_l_entete(bureau, capsys):
    _ecrire_source(bureau, [["id", "zone 1", "zone 2", "sens"]])
    resultat = format_excel_file(LOT, INSEE, verbose=False)

    assert resultat == _excel_destination_file_path(INSEE)
    assert "aucune donnee" in capsys.readouterr().out.lower().replace("é", "e")


def test_formules_signalees(bureau, capsys):
    _ecrire_source(
        bureau,
        [["id", "zone 1", "zone 2", "sens", "total"],
         [1, 2, "B", "Sud", "=B2*2"],
         [2, 1, "A", "Nord", "=B3*2"]],
    )
    format_excel_file(LOT, INSEE, verbose=True)
    assert "formules" in capsys.readouterr().out.lower()


def test_tri_stable_sur_ex_aequo(bureau):
    """A cles identiques, l'ordre d'origine du fichier est conserve."""
    _ecrire_source(
        bureau,
        [["id", "zone 1", "zone 2", "sens"],
         [10, 1, "A", "Nord"],
         [20, 1, "A", "Nord"],
         [30, 1, "A", "Nord"]],
    )
    resultat = format_excel_file(LOT, INSEE, verbose=False)
    assert [row[0] for row in _lire(resultat)] == [10, 20, 30]


def test_source_et_destination_identiques(bureau, monkeypatch, capsys):
    """Le fichier source ne doit jamais etre ecrase."""
    _ecrire_source(bureau, [["id", "zone 1", "zone 2", "sens"], [1, 2, "B", "Sud"]])
    source = _excel_source_file_path(LOT, INSEE)
    monkeypatch.setattr(excel, "_excel_destination_file_path", lambda insee: source)

    assert format_excel_file(LOT, INSEE, verbose=False) is None
    assert "identiques" in capsys.readouterr().out.lower()