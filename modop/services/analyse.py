"""Complète l'audit et génère les PPT vierges d'analyse d'une commune.

Cette étape n'a d'effet que si la case « Générer les PPT » de l'interface est
cochée. Décochée (valeur par défaut), le fichier Excel n'est pas touché : ni
la colonne ``Lien vers le .ppt``, ni le dossier ``Analyse``.

Cochée, un template PPT valide est requis : sans lui, ``PptTemplateError`` est
levée. Il s'agit d'une erreur de configuration du programme, pas d'une
anomalie propre à une commune — elle doit donc arrêter tout le traitement en
cours (le lot entier), pas seulement la commune sur laquelle elle est
détectée. C'est ``services.workflow`` qui porte cette distinction : voir
``prepare_lot_deliverables``, qui valide le template une fois, avant même de
commencer la première commune. Le template étant valide, pour chaque ligne du
fichier Excel trié identifiée par la colonne ``ID erreur`` : la colonne
``Lien vers le .ppt`` est complétée avec le lien SharePoint du PPT
correspondant (``build_ppt_sharepoint_link`` — ouverture navigateur), et le
PPT est généré localement dans le sous-dossier ``Analyse`` de la commune.

Le nommage des fichiers et le mécanisme de génération sont repris du
programme BPA (``services/ppt.py`` et la fonction ``faire_ppt()`` de
``services/traitement.py``) : un fichier ``cas_analyse_<ID erreur>.pptx`` par
ligne, les variables du template étant substituées avec les valeurs de la
ligne, en parallèle sur plusieurs threads.
"""

from __future__ import annotations

import os
import threading
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from urllib.parse import quote

from openpyxl import load_workbook

from modop.constants import COL_ID_ERR, COL_LINK, PPT_NAME_TEMPLATE
from modop.path_manager import _insee_analyse_path, _insee_ppt_sharepoint_base_url
from modop.services.excel import HEADER_SCAN_ROWS, _normalize_header
from modop.services.ppt import generate_ppt


class PptTemplateError(RuntimeError):
    """Case « Générer les PPT » cochée sans template PPT valide configuré.

    Erreur de configuration du programme, pas une anomalie propre à une
    commune : elle n'est volontairement pas rattrapée ici, ni traitée comme
    un échec de commune isolé par ``services.workflow``. Elle doit arrêter
    tout le traitement en cours (le lot entier), avec un message d'erreur
    clair — pas juste enregistrer un échec pour la commune en cours.
    """


@dataclass
class AnalysePptResult:
    """Bilan de la complétion des liens et de la génération des PPT."""

    total: int = 0
    liens_completes: int = 0
    ppt_generes: int = 0
    erreurs: int = 0
    avertissements: list[str] = field(default_factory=list)


def ppt_name_for(id_err: str) -> str:
    """Nom de fichier du PPT correspondant a une ligne (son ID erreur)."""
    return PPT_NAME_TEMPLATE.format(id=str(id_err).strip())


def valider_template_ppt(pptx_template_path: str | None) -> None:
    """Vérifie qu'un template PPT valide est configuré.

    Appelée dès que la case « Générer les PPT » est cochée, avant tout autre
    traitement : un template manquant ou introuvable est une erreur de
    configuration du programme, pas une anomalie propre à une commune.

    Raises:
        PptTemplateError: chemin vide, ou fichier introuvable.
    """
    if not (pptx_template_path and os.path.isfile(pptx_template_path)):
        raise PptTemplateError(
            f"Template PPT introuvable ({pptx_template_path or 'non configuré'}) "
            "— génération des PPT impossible. Configurez un template PPT "
            "valide, ou décochez « Générer les PPT »."
        )


def build_ppt_link(insee: str, id_err: str) -> str:
    """Chemin du PPT d'une ligne, dans le dossier Analyse de la commune."""
    return os.path.join(_insee_analyse_path(insee), ppt_name_for(id_err))

def build_ppt_sharepoint_link(insee: str, id_err: str) -> str:
    """Lien SharePoint du PPT d'une ligne (colonne « Lien vers le .ppt »).

    Contrairement à ``build_ppt_link`` qui renvoie le chemin local du fichier
    réellement généré, ce lien pointe vers la copie SharePoint et s'ouvre
    dans le navigateur (``?web=1``). C'est une URL : query string attachée
    avec ``?``, segment de nom de fichier encodé.
    """
    base = _insee_ppt_sharepoint_base_url(insee)
    return f"{base}/{quote(ppt_name_for(id_err))}?web=1"


def _find_id_and_link_columns(worksheet) -> tuple[int | None, dict[str, int]]:
    """Cherche les colonnes ``ID erreur`` et ``Lien vers le .ppt``.

    Recherche insensible a la casse et aux accents, sur les premieres lignes
    de la feuille (meme logique que ``services.excel._find_header``).

    Returns:
        (ligne d'entete, {libelle normalise: colonne}), ou (None, {}) si
        l'une des deux colonnes est absente de la feuille.
    """
    wanted = {_normalize_header(COL_ID_ERR), _normalize_header(COL_LINK)}
    max_scan = min(HEADER_SCAN_ROWS, worksheet.max_row or 1)

    for row_index in range(1, max_scan + 1):
        headers: dict[str, int] = {}
        for cell in worksheet[row_index]:
            key = _normalize_header(cell.value)
            if key and key not in headers:
                headers[key] = cell.column

        if wanted.issubset(headers):
            return row_index, headers

    return None, {}


def generer_ppts_commune(
    excel_path: str,
    insee: str,
    pptx_template_path: str | None,
    generate_ppts: bool,
    parallel: bool = True,
    verbose: bool = True,
) -> AnalysePptResult:
    """Complète la colonne ``Lien vers le .ppt`` et génère les PPT vierges.

    Sans effet si ``generate_ppts`` est faux (case décochée, valeur par
    défaut) : le fichier Excel n'est pas ouvert, la colonne lien n'est pas
    touchée.

    Coché, un template PPT valide est requis (voir ``valider_template_ppt``).
    Le template étant valide, le lien est complété et le PPT généré pour
    chaque ligne portant un ``ID erreur``.

    Args:
        excel_path: fichier Excel trié de la commune (celui que produit
            ``services.excel.format_excel_file``).
        insee: code INSEE de la commune, pour retrouver son dossier Analyse.
        pptx_template_path: chemin du template PPT vierge, configurable
            depuis l'interface. Requis si ``generate_ppts`` est vrai.
        generate_ppts: correspond a la case a cocher de l'interface. Si faux,
            rien n'est fait : ni le lien, ni le PPT.
        parallel: genere les PPT en parallele (plusieurs threads).
        verbose: affiche la progression.

    Returns:
        Le bilan de l'operation. Une commune sans les colonnes attendues, ou
        dont l'audit est absent, renvoie un bilan vide assorti d'un simple
        avertissement (anomalie de données, non bloquante — voir
        ``PptTemplateError`` pour ce qui, à l'inverse, doit bloquer).

    Raises:
        PptTemplateError: la case est cochée sans template PPT valide
            configuré. Erreur de configuration : à la charge de l'appelant
            de ne pas la traiter comme un simple échec de commune (voir
            ``services.workflow``).
    """

    def log(message: str) -> None:
        if verbose:
            print(message)

    result = AnalysePptResult()

    if not generate_ppts:
        # Case decochee : rien a faire, le fichier n'est meme pas ouvert.
        return result

    try:
        valider_template_ppt(pptx_template_path)
    except PptTemplateError as error:
        log(f"  ❌ {error}")
        raise

    if not excel_path or not os.path.isfile(excel_path):
        result.avertissements.append(f"Fichier Excel introuvable : {excel_path}")
        return result

    workbook = None
    try:
        workbook = load_workbook(excel_path)

        worksheet = header_row = headers = None
        for feuille in workbook.worksheets:
            trouve_ligne, trouve_entetes = _find_id_and_link_columns(feuille)
            if trouve_ligne is not None:
                worksheet, header_row, headers = feuille, trouve_ligne, trouve_entetes
                break

        if worksheet is None:
            message = (
                f"Colonnes '{COL_ID_ERR}' / '{COL_LINK}' introuvables : "
                "liens PPT non complétés."
            )
            result.avertissements.append(message)
            log(f"  ⚠️ {message}")
            return result

        col_id = headers[_normalize_header(COL_ID_ERR)]
        col_link = headers[_normalize_header(COL_LINK)]

        # Toutes les colonnes de l'entete, pour la substitution des variables
        # du template (VARIABLES associe chaque variable a un nom de colonne).
        col_names = {
            c: str(worksheet.cell(header_row, c).value).strip()
            for c in range(1, worksheet.max_column + 1)
            if worksheet.cell(header_row, c).value
        }

        # id_err -> (chemin du PPT, donnees de la ligne)
        ppt_plan: dict[str, tuple[str, dict]] = {}

        for row_index in range(header_row + 1, worksheet.max_row + 1):
            id_err = str(worksheet.cell(row_index, col_id).value or "").strip()
            if not id_err:
                continue

            result.total += 1
            dst = build_ppt_link(insee, id_err)  # chemin local, pour generate_ppt
            # La cellule reçoit le lien SharePoint (ouverture navigateur), pas
            # le chemin local : c'est ce que l'utilisateur clique dans l'audit.
            worksheet.cell(row_index, col_link).value = build_ppt_sharepoint_link(
                insee, id_err
            )
            result.liens_completes += 1

            row_data = {
                nom: worksheet.cell(row_index, c).value
                for c, nom in col_names.items()
            }
            ppt_plan[id_err] = (dst, row_data)

        if ppt_plan:
            os.makedirs(_insee_analyse_path(insee), exist_ok=True)
            lock = threading.Lock()

            def faire_ppt(item: tuple[str, tuple[str, dict]]) -> None:
                id_err, (dst, row_data) = item
                try:
                    generate_ppt(row_data, dst, pptx_template_path)
                    with lock:
                        result.ppt_generes += 1
                except Exception as error:
                    with lock:
                        result.erreurs += 1
                    log(f"  ❌ PPT '{os.path.basename(dst)}' : "
                        f"{type(error).__name__} : {error}")

            items = list(ppt_plan.items())
            if parallel and len(items) > 1:
                with ThreadPoolExecutor(
                    max_workers=min(8, (os.cpu_count() or 4) * 2)
                ) as executor:
                    list(executor.map(faire_ppt, items))
            else:
                for item in items:
                    faire_ppt(item)

        workbook.save(excel_path)

        log(f"  🖼️  Liens PPT : {result.liens_completes}/{result.total}"
            f"  Générés : {result.ppt_generes}"
            + (f"  Erreurs : {result.erreurs}" if result.erreurs else ""))

        return result

    except Exception as error:
        message = f"Erreur génération PPT : {type(error).__name__} : {error}"
        result.avertissements.append(message)
        log(f"  ❌ {message}")
        return result

    finally:
        if workbook is not None:
            workbook.close()


if __name__ == "__main__":
    # uv run python -m modop.services.analyse
    from modop.path_manager import _excel_destination_file_path, pptx_template_path
    from modop.services.config_io import load_and_apply

    # Les chemins et la case "Générer les PPT" enregistrés depuis l'interface
    # s'appliquent aussi en console.
    _config = load_and_apply()

    INSEE = "45001"
    resultat = generer_ppts_commune(
        _excel_destination_file_path(INSEE),
        INSEE,
        pptx_template_path(),
        generate_ppts=_config.get("generate_ppts", False),
    )
    print(resultat)
